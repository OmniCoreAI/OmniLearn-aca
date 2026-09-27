"""Bulk member import from Excel / CSV into an entity.

1. ``validate_upload`` parses the file (xlsx via openpyxl in read-only mode, or
   CSV), checks every row and stores an ``ImportJob`` with ``ImportRow``s.
2. The reviewer commits: rows are processed one by one (a failing row never
   stops the others) — new accounts are created, existing ones linked, members
   added to their groups, then audiences are re-synced and new accounts get a
   "set your password" message.
3. Failed / invalid rows can be downloaded as an xlsx with the reason.

Academy staff and coordinators with ``can_import_users`` may import.
"""
import asyncio
import csv
import io
import logging
import re
import zipfile
from datetime import datetime
from typing import Dict, List, Optional, Tuple
from uuid import uuid4

from fastapi import HTTPException, UploadFile
from sqlmodel import func, select
from sqlmodel.ext.asyncio.session import AsyncSession

from src.db.administration.entities import Entity, EntityMember, Position
from src.db.administration.imports import (
    ImportCommitOptions,
    ImportJob,
    ImportJobRead,
    ImportRow,
    ImportRowPage,
    ImportRowRead,
)
from src.db.administration.lookups import ConfigStatus
from src.db.user_organizations import UserOrganization
from src.db.usergroups import UserGroup, UserGroupType
from src.db.users import User
from src.security.rbac.constants import TRAINEE_ROLE_ID
from src.services.administration.authz import AnyUser
from src.services.administration.common import bad_request, get_by_uuid_or_404, now
from src.services.administration.entities import (
    get_entity_by_uuid,
    require_entity_access,
    upsert_member,
)
from src.services.notifications.sms import normalize_phone

logger = logging.getLogger(__name__)

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
FIELDS = ["first_name", "last_name", "email", "phone", "employee_id", "position", "groups"]
HEADERS = {
    "first_name": "First name / الاسم الأول",
    "last_name": "Last name / اسم العائلة",
    "email": "Email / البريد الإلكتروني",
    "phone": "Phone / الهاتف",
    "employee_id": "Employee ID / الرقم الوظيفي",
    "position": "Position / المنصب",
    "groups": "Groups / المجموعات",
}
ALIASES = {
    "first_name": {"first name", "firstname", "first", "name", "الاسم الأول", "الاسم الاول", "الاسم"},
    "last_name": {"last name", "lastname", "last", "surname", "family name", "اسم العائلة", "اللقب"},
    "email": {"email", "e-mail", "email address", "mail", "البريد الإلكتروني", "البريد الالكتروني", "البريد"},
    "phone": {"phone", "mobile", "phone number", "mobile number", "الهاتف", "الموبايل", "رقم الهاتف", "الجوال"},
    "employee_id": {"employee id", "employee_id", "employee no", "staff id", "id", "الرقم الوظيفي", "رقم الموظف"},
    "position": {"position", "job title", "title", "المنصب", "الوظيفة", "المسمى الوظيفي"},
    "groups": {"groups", "group", "department", "المجموعات", "المجموعة", "الإدارة", "القسم"},
}
GROUP_SPLIT_RE = re.compile(r"[;،,|]")


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


def _header_key(value: object) -> Optional[str]:
    text = " ".join(str(value or "").strip().lower().replace("_", " ").split())
    if not text:
        return None
    for field, headers in HEADERS.items():
        if text == headers.lower():
            return field
    for part in [text, *[p.strip() for p in text.split("/")]]:
        for field, aliases in ALIASES.items():
            if part in aliases:
                return field
    return None


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _rows_from_xlsx(data: bytes) -> List[List[str]]:
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if any(name.lower().endswith("vbaproject.bin") for name in archive.namelist()):
                raise bad_request("Macro-enabled workbooks are not accepted")
    except zipfile.BadZipFile:
        raise bad_request("The file is not a valid .xlsx workbook")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise bad_request("The workbook could not be read")
    try:
        sheet = workbook.worksheets[0]
        rows: List[List[str]] = []
        for values in sheet.iter_rows(values_only=True):
            rows.append([_cell(v) for v in values])
            if len(rows) > MAX_ROWS + 20:
                break
        return rows
    finally:
        workbook.close()


def _rows_from_csv(data: bytes) -> List[List[str]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1256", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = []
    for row in reader:
        rows.append([_cell(v) for v in row])
        if len(rows) > MAX_ROWS + 20:
            break
    return rows


def parse_rows(file_name: str, data: bytes) -> List[Tuple[int, Dict[str, str]]]:
    """Return ``[(spreadsheet_row_number, {field: value})]`` for non-empty rows."""
    name = (file_name or "").lower()
    if name.endswith((".xlsm", ".xltm", ".xls")):
        raise bad_request("Upload an .xlsx or .csv file")
    if len(data) > MAX_BYTES:
        raise bad_request("The file is larger than 5 MB")
    if name.endswith(".csv"):
        raw = _rows_from_csv(data)
    elif name.endswith(".xlsx") or data[:2] == b"PK":
        raw = _rows_from_xlsx(data)
    else:
        raise bad_request("Upload an .xlsx or .csv file")

    header_index = next((i for i, r in enumerate(raw) if any(r)), None)
    if header_index is None:
        raise bad_request("The file is empty")
    columns = {i: _header_key(v) for i, v in enumerate(raw[header_index])}
    if "email" not in columns.values() or "first_name" not in columns.values():
        raise bad_request("Columns 'First name' and 'Email' are required — download the template")
    out = []
    for offset, values in enumerate(raw[header_index + 1:], start=header_index + 2):
        if not any(values):
            continue
        record = {field: "" for field in FIELDS}
        for i, value in enumerate(values):
            field = columns.get(i)
            if field and value:
                record[field] = value
        out.append((offset, record))
    if len(out) > MAX_ROWS:
        raise bad_request(f"At most {MAX_ROWS} rows can be imported at once")
    if not out:
        raise bad_request("No rows to import")
    return out


# ---------------------------------------------------------------------------
# Lookups for validation
# ---------------------------------------------------------------------------


async def _entity_positions(db_session: AsyncSession, entity: Entity) -> Dict[str, Position]:
    rows = (
        await db_session.execute(
            select(Position).where(
                Position.org_id == entity.org_id,
                Position.status == ConfigStatus.ACTIVE.value,
                (Position.entity_id == entity.id) | (Position.entity_id.is_(None)),  # type: ignore[union-attr]
            )
        )
    ).scalars().all()
    index: Dict[str, Position] = {}
    for p in rows:
        for key in (p.name, p.code):
            if key:
                # Entity-specific positions win over shared ones with the same name.
                if p.entity_id or key.lower() not in index:
                    index[key.strip().lower()] = p
    return index


async def _entity_groups(db_session: AsyncSession, entity: Entity) -> Dict[str, UserGroup]:
    rows = (
        await db_session.execute(
            select(UserGroup).where(
                UserGroup.entity_id == entity.id,
                UserGroup.group_type != UserGroupType.SYSTEM.value,
                UserGroup.managed_key.is_(None),  # type: ignore[union-attr]
                UserGroup.status != "inactive",
            )
        )
    ).scalars().all()
    return {g.name.strip().lower(): g for g in rows}


# ---------------------------------------------------------------------------
# Reads
# ---------------------------------------------------------------------------


async def _job_read(db_session: AsyncSession, job: ImportJob) -> ImportJobRead:
    entity = await db_session.get(Entity, job.entity_id)
    counts = dict(
        (
            await db_session.execute(
                select(ImportRow.status, func.count(ImportRow.id)).where(ImportRow.job_id == job.id).group_by(ImportRow.status)
            )
        ).all()
    )
    return ImportJobRead(
        **job.model_dump(exclude={"id", "org_id", "entity_id", "created_by_user_id", "update_date"}),
        entity_uuid=entity.entity_uuid if entity else "",
        entity_name=entity.name if entity else "",
        counts={k: int(v) for k, v in counts.items()},
    )


async def _job_with_access(
    db_session: AsyncSession, current_user: AnyUser, job_uuid: str, action: str = "read"
) -> Tuple[ImportJob, Entity]:
    job = await get_by_uuid_or_404(db_session, ImportJob, ImportJob.job_uuid, job_uuid, "Import")
    entity = await db_session.get(Entity, job.entity_id)
    if entity is None:
        raise HTTPException(status_code=404, detail="Entity not found")
    await require_entity_access(db_session, current_user, entity, "can_import_users", action)  # type: ignore[arg-type]
    return job, entity


async def list_jobs(db_session: AsyncSession, current_user: AnyUser, entity_uuid: str) -> List[ImportJobRead]:
    entity = await get_entity_by_uuid(db_session, entity_uuid)
    await require_entity_access(db_session, current_user, entity, "can_import_users")
    jobs = (
        await db_session.execute(
            select(ImportJob).where(ImportJob.entity_id == entity.id).order_by(ImportJob.id.desc()).limit(50)  # type: ignore[union-attr]
        )
    ).scalars().all()
    return [await _job_read(db_session, j) for j in jobs]


async def get_job(db_session: AsyncSession, current_user: AnyUser, job_uuid: str) -> ImportJobRead:
    job, _entity = await _job_with_access(db_session, current_user, job_uuid)
    return await _job_read(db_session, job)


async def list_rows(
    db_session: AsyncSession,
    current_user: AnyUser,
    job_uuid: str,
    status: Optional[str] = None,
    page: int = 1,
    limit: int = 100,
) -> ImportRowPage:
    job, _entity = await _job_with_access(db_session, current_user, job_uuid)
    page, limit = max(1, page), max(1, min(limit, 500))
    stmt = select(ImportRow).where(ImportRow.job_id == job.id)
    if status:
        stmt = stmt.where(ImportRow.status.in_(status.split(",")))  # type: ignore[attr-defined]
    total = (await db_session.execute(select(func.count()).select_from(stmt.subquery()))).scalar() or 0
    rows = (
        await db_session.execute(stmt.order_by(ImportRow.row_number).offset((page - 1) * limit).limit(limit))
    ).scalars().all()
    return ImportRowPage(
        items=[ImportRowRead(row_number=r.row_number, data=r.data, status=r.status, errors=r.errors or []) for r in rows],
        total=int(total),
        page=page,
        limit=limit,
    )


# ---------------------------------------------------------------------------
# Validate
# ---------------------------------------------------------------------------


async def validate_upload(
    db_session: AsyncSession, current_user: AnyUser, entity_uuid: str, upload: UploadFile
) -> ImportJobRead:
    entity = await get_entity_by_uuid(db_session, entity_uuid)
    access = await require_entity_access(db_session, current_user, entity, "can_import_users", "create")
    data = await upload.read(MAX_BYTES + 1)
    rows = parse_rows(upload.filename or "", data)

    positions = await _entity_positions(db_session, entity)
    groups = await _entity_groups(db_session, entity)
    emails = [r["email"].strip().lower() for _, r in rows if r.get("email")]
    users_by_email: Dict[str, User] = {}
    for chunk_start in range(0, len(emails), 500):
        chunk = emails[chunk_start:chunk_start + 500]
        for u in (await db_session.execute(select(User).where(func.lower(User.email).in_(chunk)))).scalars().all():
            users_by_email[(u.email or "").lower()] = u
    org_member_ids = set(
        (
            await db_session.execute(
                select(UserOrganization.user_id).where(
                    UserOrganization.org_id == entity.org_id,
                    UserOrganization.user_id.in_([u.id for u in users_by_email.values()] or [0]),  # type: ignore[attr-defined]
                )
            )
        ).scalars().all()
    )
    taken_employee_ids = dict(
        (
            await db_session.execute(
                select(EntityMember.employee_id, EntityMember.user_id).where(
                    EntityMember.entity_id == entity.id, EntityMember.employee_id.is_not(None)  # type: ignore[union-attr]
                )
            )
        ).all()
    )

    job = ImportJob(
        org_id=entity.org_id,
        entity_id=entity.id,  # type: ignore[arg-type]
        created_by_user_id=access.user_id,
        file_name=(upload.filename or "import")[:200],
        status="validated",
        total_rows=len(rows),
        job_uuid=f"import_{uuid4()}",
        creation_date=now(),
        update_date=now(),
    )
    db_session.add(job)
    await db_session.flush()

    seen_emails: Dict[str, int] = {}
    seen_employee: Dict[str, int] = {}
    valid = invalid = 0
    for row_number, record in rows:
        errors: List[str] = []
        clean = {k: (v or "").strip() for k, v in record.items()}
        email = clean["email"].lower()
        clean["email"] = email
        if not clean["first_name"]:
            errors.append("First name is required")
        if not email:
            errors.append("Email is required")
        elif not EMAIL_RE.match(email):
            errors.append("Invalid email address")
        elif email in seen_emails:
            errors.append(f"Duplicate of row {seen_emails[email]}")
        else:
            seen_emails[email] = row_number
        if clean["phone"]:
            phone = normalize_phone(clean["phone"])
            if phone is None:
                errors.append("Invalid phone number")
            else:
                clean["phone"] = phone
        existing = users_by_email.get(email)
        if clean["employee_id"]:
            emp = clean["employee_id"]
            if emp in seen_employee:
                errors.append(f"Employee ID repeated from row {seen_employee[emp]}")
            else:
                seen_employee[emp] = row_number
            owner = taken_employee_ids.get(emp)
            if owner is not None and (existing is None or owner != existing.id):
                errors.append("Employee ID already used by another member of this entity")
        if clean["position"]:
            position = positions.get(clean["position"].lower())
            if position is None:
                errors.append(f"Unknown position: {clean['position']}")
            else:
                clean["position_uuid"] = position.position_uuid
        group_uuids = []
        for name in [g.strip() for g in GROUP_SPLIT_RE.split(clean["groups"]) if g.strip()]:
            group = groups.get(name.lower())
            if group is None:
                errors.append(f"Unknown group: {name}")
            else:
                group_uuids.append(group.usergroup_uuid)
        clean["group_uuids"] = group_uuids
        status = "valid"
        if existing is not None:
            if existing.id in org_member_ids:
                status = "existing"
            elif not access.is_academy:
                errors.append("This email already has an account outside the academy — ask the academy to add it")
            else:
                status = "existing"
        if errors:
            status = "invalid"
            invalid += 1
        else:
            valid += 1
        db_session.add(ImportRow(job_id=job.id, row_number=row_number, data=clean, status=status, errors=errors))  # type: ignore[arg-type]
    job.valid_rows, job.invalid_rows = valid, invalid
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)
    return await _job_read(db_session, job)


# ---------------------------------------------------------------------------
# Commit
# ---------------------------------------------------------------------------

_background_tasks: set = set()


async def commit_job(
    db_session: AsyncSession, current_user: AnyUser, job_uuid: str, options: ImportCommitOptions
) -> ImportJobRead:
    job, entity = await _job_with_access(db_session, current_user, job_uuid, "create")
    if job.status != "validated":
        raise HTTPException(status_code=409, detail="This import was already started")
    if not job.valid_rows:
        raise bad_request("There are no valid rows to import")
    if options.group_uuid:
        group = await get_by_uuid_or_404(db_session, UserGroup, UserGroup.usergroup_uuid, options.group_uuid, "Group")
        if group.entity_id != entity.id or group.managed_key or group.group_type == UserGroupType.SYSTEM.value:
            raise bad_request("Choose one of this entity's groups")
    if options.resource_uuid:
        from src.db.administration.audience import AudienceAssignment, AudienceMode, AudienceType
        from src.services.administration.authz import has_admin_permission

        acting = getattr(current_user, "id", 0)
        if not await has_admin_permission(db_session, acting, entity.org_id, "entities", "create"):
            available = (
                await db_session.execute(
                    select(AudienceAssignment.id).where(
                        AudienceAssignment.resource_type == (options.resource_type or "course"),
                        AudienceAssignment.resource_uuid == options.resource_uuid,
                        AudienceAssignment.audience_type == AudienceType.ENTITY.value,
                        AudienceAssignment.audience_id == entity.id,
                        AudienceAssignment.mode == AudienceMode.AVAILABLE.value,
                    )
                )
            ).first()
            if not available:
                raise HTTPException(status_code=403, detail="The academy has not made this available to your entity")
    job.status = "processing"
    job.options = options.model_dump()
    job.update_date = now()
    db_session.add(job)
    await db_session.commit()

    import os

    if os.environ.get("TESTING") == "true":
        await process_job(db_session, job.id, getattr(current_user, "id", None))  # type: ignore[arg-type]
    else:
        task = asyncio.create_task(_process_in_background(job.id, getattr(current_user, "id", None)))  # type: ignore[arg-type]
        _background_tasks.add(task)
        task.add_done_callback(_background_tasks.discard)
    await db_session.refresh(job)
    return await _job_read(db_session, job)


async def _process_in_background(job_id: int, acting_user_id: Optional[int]) -> None:
    from src.core.events.database import _async_session_factory

    async with _async_session_factory() as session:
        try:
            await process_job(session, job_id, acting_user_id)
        except Exception as exc:
            logger.exception("Import %s failed", job_id)
            await session.rollback()
            job = await session.get(ImportJob, job_id)
            if job is not None:
                job.status = "failed"
                job.error = str(exc)[:500]
                session.add(job)
                await session.commit()


async def process_job(db_session: AsyncSession, job_id: int, acting_user_id: Optional[int]) -> None:
    from src.services.orgs.users import provision_org_user

    job = await db_session.get(ImportJob, job_id)
    if job is None:
        return
    entity = await db_session.get(Entity, job.entity_id)
    if entity is None:
        return
    options = ImportCommitOptions(**(job.options or {}))
    org_id, entity_id = entity.org_id, entity.id
    extra_group_id: Optional[int] = None
    if options.group_uuid:
        extra_group_id = (
            await db_session.execute(select(UserGroup.id).where(UserGroup.usergroup_uuid == options.group_uuid))
        ).scalars().first()
    # Plain ids/values only: a failed row rolls back and expires ORM objects.
    row_ids = list(
        (
            await db_session.execute(
                select(ImportRow.id).where(ImportRow.job_id == job.id, ImportRow.status.in_(["valid", "existing"])).order_by(ImportRow.row_number)  # type: ignore[attr-defined]
            )
        ).scalars().all()
    )
    group_ids_by_uuid = {g.usergroup_uuid: g.id for g in (await _entity_groups(db_session, entity)).values()}
    position_ids_by_uuid = {p.position_uuid: p.id for p in (await _entity_positions(db_session, entity)).values()}

    created_ids: List[int] = []
    imported_ids: List[int] = []
    created = existing = failed = 0
    for row_id in row_ids:
        row = await db_session.get(ImportRow, row_id)
        data = dict(row.data or {}) if row else {}
        try:
            user = (
                await db_session.execute(select(User).where(func.lower(User.email) == data.get("email", "")))
            ).scalars().first()
            is_new = user is None
            if is_new:
                user, _password = await provision_org_user(
                    db_session,
                    org_id,
                    email=data["email"],
                    first_name=data.get("first_name", ""),
                    last_name=data.get("last_name", ""),
                    role_id=TRAINEE_ROLE_ID,
                    extra_metadata={"phone": data["phone"]} if data.get("phone") else None,
                    signup_method="entity_import",
                )
            else:
                membership = (
                    await db_session.execute(
                        select(UserOrganization).where(
                            UserOrganization.user_id == user.id, UserOrganization.org_id == org_id
                        )
                    )
                ).scalars().first()
                if membership is None:
                    db_session.add(
                        UserOrganization(
                            user_id=user.id, org_id=org_id, role_id=TRAINEE_ROLE_ID,
                            creation_date=now(), update_date=now(),
                        )
                    )
                if data.get("phone") and not (user.extra_metadata or {}).get("phone"):
                    user.extra_metadata = {**(user.extra_metadata or {}), "phone": data["phone"]}
                    db_session.add(user)
            group_ids = [group_ids_by_uuid[g] for g in data.get("group_uuids", []) if g in group_ids_by_uuid]
            if extra_group_id is not None:
                group_ids.append(extra_group_id)
            await upsert_member(
                db_session,
                entity,
                user,  # type: ignore[arg-type]
                position_id=position_ids_by_uuid.get(data.get("position_uuid") or ""),
                employee_id=data.get("employee_id") or None,
                group_ids=group_ids,
            )
            row = await db_session.get(ImportRow, row_id)
            row.status = "created" if is_new else "linked"  # type: ignore[union-attr]
            row.user_id = user.id  # type: ignore[union-attr]
            db_session.add(row)
            await db_session.commit()
            imported_ids.append(user.id)  # type: ignore[union-attr]
            if is_new:
                created += 1
                created_ids.append(user.id)  # type: ignore[union-attr]
            else:
                existing += 1
        except Exception as exc:
            await db_session.rollback()
            await db_session.refresh(entity)
            failed += 1
            row = await db_session.get(ImportRow, row_id)
            if row is not None:
                row.status = "failed"
                detail = getattr(exc, "detail", None) or str(exc) or exc.__class__.__name__
                row.errors = [str(detail)[:300]]
                db_session.add(row)
                await db_session.commit()

    # Training: audiences that include this entity pick up the new members.
    from src.services.administration.audience import resync_affected

    await resync_affected(db_session, org_id, entity_ids=[entity_id], user_ids=imported_ids)  # type: ignore[list-item]
    if extra_group_id is not None:
        await resync_affected(db_session, org_id, group_ids=[extra_group_id])
    if options.resource_uuid and imported_ids:
        entity = await db_session.get(Entity, entity_id)
        await _assign_resource(db_session, entity, options, imported_ids, acting_user_id)  # type: ignore[arg-type]
    if options.notify and created_ids:
        from src.services.notifications.events import password_setup

        for user_id in created_ids:
            await password_setup(db_session, org_id, user_id)

    job = await db_session.get(ImportJob, job_id)
    job.status = "completed"  # type: ignore[union-attr]
    job.created_count, job.existing_count, job.failed_count = created, existing, failed  # type: ignore[union-attr]
    job.completed_at = now()  # type: ignore[union-attr]
    job.update_date = now()  # type: ignore[union-attr]
    db_session.add(job)
    await db_session.commit()


async def _assign_resource(
    db_session: AsyncSession, entity: Entity, options: ImportCommitOptions, user_ids: List[int], acting_user_id: Optional[int]
) -> None:
    """Assign the chosen course / program to each imported member."""
    from src.db.administration.audience import AudienceAssignment, AudienceMode, AudienceType
    from src.services.administration.audience import sync_resource_audience

    resource_type = options.resource_type or "course"
    existing = set(
        (
            await db_session.execute(
                select(AudienceAssignment.audience_id).where(
                    AudienceAssignment.resource_type == resource_type,
                    AudienceAssignment.resource_uuid == options.resource_uuid,
                    AudienceAssignment.audience_type == AudienceType.USER.value,
                    AudienceAssignment.entity_id == entity.id,
                )
            )
        ).scalars().all()
    )
    for user_id in user_ids:
        if user_id in existing:
            continue
        db_session.add(
            AudienceAssignment(
                org_id=entity.org_id,
                resource_type=resource_type,
                resource_uuid=options.resource_uuid,  # type: ignore[arg-type]
                audience_type=AudienceType.USER.value,
                audience_id=user_id,
                entity_id=entity.id,
                mode=AudienceMode.ASSIGNED,
                notify=True,
                created_by_user_id=acting_user_id,
                assignment_uuid=f"audience_{uuid4()}",
                creation_date=now(),
                update_date=now(),
            )
        )
    await db_session.commit()
    await sync_resource_audience(db_session, resource_type, options.resource_uuid)  # type: ignore[arg-type]


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------


def _safe(value: object) -> object:
    from src.services.orgs.users import _csv_safe

    return _csv_safe(value)


async def template_workbook(db_session: AsyncSession, current_user: AnyUser, entity_uuid: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    entity = await get_entity_by_uuid(db_session, entity_uuid)
    await require_entity_access(db_session, current_user, entity, "can_import_users")
    positions = sorted({p.name for p in (await _entity_positions(db_session, entity)).values()})
    groups = sorted({g.name for g in (await _entity_groups(db_session, entity)).values()})

    wb = Workbook()
    ws = wb.active
    ws.title = "Members"
    ws.sheet_view.rightToLeft = False
    for col, field in enumerate(FIELDS, start=1):
        cell = ws.cell(row=1, column=col, value=HEADERS[field])
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        ws.column_dimensions[cell.column_letter].width = 26
    ws.freeze_panes = "A2"

    lists = wb.create_sheet("Lists")
    lists.cell(row=1, column=1, value="Positions")
    lists.cell(row=1, column=2, value="Groups")
    for i, name in enumerate(positions, start=2):
        lists.cell(row=i, column=1, value=_safe(name))
    for i, name in enumerate(groups, start=2):
        lists.cell(row=i, column=2, value=_safe(name))
    lists.sheet_state = "hidden"
    if positions:
        dv = DataValidation(type="list", formula1=f"=Lists!$A$2:$A${len(positions) + 1}", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"F2:F{MAX_ROWS + 1}")
    if groups:
        # One group per cell from the list; type several separated by ";" if needed.
        dv = DataValidation(type="list", formula1=f"=Lists!$B$2:$B${len(groups) + 1}", allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(f"G2:G{MAX_ROWS + 1}")

    notes = wb.create_sheet("Instructions")
    for i, line in enumerate(
        [
            "Fill one member per row on the Members sheet. First name and Email are required.",
            "املأ عضواً واحداً في كل صف في ورقة Members. الاسم الأول والبريد الإلكتروني مطلوبان.",
            "Phone: Egyptian numbers like 01012345678 are accepted.",
            "Position and Groups must match the names defined for the entity; separate several groups with ;",
            "المنصب والمجموعات يجب أن تطابق الأسماء المعرفة للجهة؛ افصل بين المجموعات بـ ;",
            f"Entity: {entity.name}",
        ],
        start=1,
    ):
        notes.cell(row=i, column=1, value=_safe(line))
    notes.column_dimensions["A"].width = 110

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def failed_rows_workbook(db_session: AsyncSession, current_user: AnyUser, job_uuid: str) -> Tuple[str, bytes]:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    job, _entity = await _job_with_access(db_session, current_user, job_uuid)
    rows = (
        await db_session.execute(
            select(ImportRow).where(ImportRow.job_id == job.id, ImportRow.status.in_(["invalid", "failed"])).order_by(ImportRow.row_number)  # type: ignore[attr-defined]
        )
    ).scalars().all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Failed rows"
    header = ["Row", *[HEADERS[f] for f in FIELDS], "Error / الخطأ"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        data = r.data or {}
        ws.append([r.row_number, *[_safe(data.get(f, "")) for f in FIELDS], _safe("; ".join(r.errors or []))])
    buffer = io.BytesIO()
    wb.save(buffer)
    stamp = datetime.now().strftime("%Y%m%d")
    return f"failed-rows-{stamp}.xlsx", buffer.getvalue()
