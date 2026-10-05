"""Entity member import (xlsx / csv), coordinator progress and instructor invitations."""
import io
from contextlib import ExitStack
from datetime import datetime
from unittest.mock import AsyncMock, patch

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from sqlmodel import select

from src.db.administration.audience import AudienceAssignmentCreate
from src.db.administration.entities import (
    CoordinatorAssign,
    CoordinatorPermissions,
    EntityCreate,
    EntityGroupCreate,
    EntityInstructorInvite,
    EntityMemberCreate,
    EntityUpdate,
    PositionCreate,
)
from src.db.administration.imports import ImportCommitOptions
from src.db.courses.courses import Course
from src.db.instructors.instructors import Instructor
from src.db.trail_runs import TrailRun
from src.db.trails import Trail
from src.db.user_organizations import UserOrganization
from src.db.users import PublicUser, User
from src.services.administration import audience as aud_svc
from src.services.administration import entities as ent_svc
from src.services.administration import entity_coordination as coord_svc
from src.services.administration import imports as imp_svc

HEADER = ["First name", "Last name", "Email", "Phone", "Employee ID", "Position", "Groups"]


def _patches():
    stack = ExitStack()
    for target in (
        "src.services.orgs.users.check_limits_with_usage",
        "src.services.orgs.users.increase_feature_usage",
        "src.services.orgs.users.track",
        "src.services.orgs.users.dispatch_webhooks",
        "src.services.admin.admin.track",
    ):
        stack.enter_context(patch(target, new_callable=AsyncMock))
    stack.enter_context(patch("src.services.users.password_reset.create_password_setup_code", return_value="Abc12345"))
    return stack


def _xlsx(rows, header=HEADER) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(data: bytes, name="members.xlsx") -> UploadFile:
    return UploadFile(file=io.BytesIO(data), filename=name)


async def _user(db, org, uid, username, role_id=4) -> PublicUser:
    db.add(User(id=uid, username=username, first_name=username, last_name="T", email=f"{username}@test.com", password="x",
                user_uuid=f"user_{username}", creation_date="", update_date=""))
    await db.commit()
    db.add(UserOrganization(user_id=uid, org_id=org.id, role_id=role_id, creation_date="", update_date=""))
    await db.commit()
    return PublicUser(id=uid, username=username, first_name=username, last_name="T", email=f"{username}@test.com", user_uuid=f"user_{username}")


async def _setup(db, org, admin_user):
    entity = await ent_svc.create_entity(db, admin_user, org.id, EntityCreate(name="Ministry"))
    await ent_svc.create_position(db, admin_user, org.id, PositionCreate(name="Engineer"))
    group = await ent_svc.create_entity_group(db, admin_user, entity.entity_uuid, EntityGroupCreate(name="IT Department"))
    return entity, group


class TestParsing:
    def test_aliases_arabic_headers_and_limits(self):
        rows = imp_svc.parse_rows("m.xlsx", _xlsx([["Ali", "", "ALI@x.com", 1012345678.0, "", "", ""]], header=["الاسم الأول", "Family name", "البريد الإلكتروني", "الهاتف", "", "", ""]))
        assert rows == [(2, {"first_name": "Ali", "last_name": "", "email": "ALI@x.com", "phone": "1012345678", "employee_id": "", "position": "", "groups": ""})]
        csv_rows = imp_svc.parse_rows("m.csv", "First name,Email\nMona,mona@x.com\n,\n".encode("utf-8-sig"))
        assert [r[1]["email"] for r in csv_rows] == ["mona@x.com"]
        for name, data, msg in (
            ("m.xlsm", b"x", "xlsx"),
            ("m.xlsx", b"not a zip", "valid"),
            ("m.xlsx", _xlsx([["a", "b"]], header=["Name only", "Other"]), "required"),
            ("m.pdf", b"%PDF", "xlsx"),
        ):
            with pytest.raises(HTTPException) as exc:
                imp_svc.parse_rows(name, data)
            assert msg in exc.value.detail

    def test_macro_workbook_rejected(self):
        import zipfile

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("xl/vbaProject.bin", b"macro")
        with pytest.raises(HTTPException) as exc:
            imp_svc.parse_rows("m.xlsx", buf.getvalue())
        assert "Macro" in exc.value.detail


class TestImport:
    @pytest.mark.asyncio
    async def test_validate_review_commit_and_failed_rows(self, db, org, admin_user, regular_user):
        entity, group = await _setup(db, org, admin_user)
        data = _xlsx([
            ["Mona", "Ali", "mona@ministry.gov.eg", "01012345678", "E1", "Engineer", "IT Department"],
            ["Omar", "", "omar@ministry.gov.eg", "", "E2", "", ""],
            ["", "", "bad-email", "", "", "", ""],                       # invalid: no name, bad email
            ["Dup", "", "MONA@ministry.gov.eg", "", "", "", ""],          # duplicate of row 2
            ["Reg", "", "regular@test.com", "", "E1", "", ""],            # existing account, employee id repeated
            ["Pos", "", "pos@ministry.gov.eg", "", "", "Astronaut", "Nope"],
            ["Regular", "", "regular@test.com", "", "E9", "", ""],        # duplicate email (row 6)
        ])
        job = await imp_svc.validate_upload(db, admin_user, entity.entity_uuid, _upload(data))
        assert job.total_rows == 7 and job.valid_rows == 2 and job.invalid_rows == 5
        rows = {r.row_number: r for r in (await imp_svc.list_rows(db, admin_user, job.job_uuid)).items}
        assert rows[2].status == "valid" and rows[2].data["phone"] == "+201012345678"
        assert "Invalid email address" in rows[4].errors and "First name is required" in rows[4].errors
        assert rows[5].errors == ["Duplicate of row 2"]
        assert any("Employee ID repeated" in e for e in rows[6].errors)
        assert set(rows[7].errors) == {"Unknown position: Astronaut", "Unknown group: Nope"}

        with _patches():
            done = await imp_svc.commit_job(db, admin_user, job.job_uuid, ImportCommitOptions(notify=True))
        assert done.status == "completed" and done.created_count == 2 and done.failed_count == 0
        with pytest.raises(HTTPException) as exc:
            await imp_svc.commit_job(db, admin_user, job.job_uuid, ImportCommitOptions())
        assert exc.value.status_code == 409

        members = await ent_svc.list_members(db, admin_user, entity.entity_uuid)
        assert members.total == 2
        mona = next(m for m in members.items if m.email == "mona@ministry.gov.eg")
        assert mona.position_name == "Engineer" and [g.name for g in mona.groups] == ["IT Department"]
        user = (await db.execute(select(User).where(User.email == "mona@ministry.gov.eg"))).scalar_one()
        assert user.must_change_password is True and user.extra_metadata == {"phone": "+201012345678"}

        name, content = await imp_svc.failed_rows_workbook(db, admin_user, job.job_uuid)
        sheet = load_workbook(io.BytesIO(content)).active
        values = [[c.value for c in row] for row in sheet.iter_rows()]
        assert len(values) == 6 and values[1][0] == 4 and "Invalid email" in values[1][-1]

    @pytest.mark.asyncio
    async def test_existing_members_linked_and_formula_cells_neutralized(self, db, org, admin_user, regular_user):
        entity, _group = await _setup(db, org, admin_user)
        job = await imp_svc.validate_upload(
            db, admin_user, entity.entity_uuid,
            # CSV keeps "=..." as literal text (xlsx would hold an uncalculated formula).
            _upload('First name,Email\nReg,regular@test.com\n"=HYPERLINK(""x"")",not-an-email\n'.encode(), "m.csv"),
        )
        rows = (await imp_svc.list_rows(db, admin_user, job.job_uuid, status="existing")).items
        assert [r.data["email"] for r in rows] == ["regular@test.com"]
        with _patches():
            done = await imp_svc.commit_job(db, admin_user, job.job_uuid, ImportCommitOptions(notify=False))
        assert done.existing_count == 1 and done.created_count == 0
        _name, content = await imp_svc.failed_rows_workbook(db, admin_user, job.job_uuid)
        cell = load_workbook(io.BytesIO(content)).active.cell(row=2, column=2).value
        assert cell.startswith("'=")

    @pytest.mark.asyncio
    async def test_coordinator_scope(self, db, org, admin_user):
        entity, _group = await _setup(db, org, admin_user)
        other = await ent_svc.create_entity(db, admin_user, org.id, EntityCreate(name="University"))
        coord = await _user(db, org, 30, "coord")
        await ent_svc.assign_coordinator(db, admin_user, entity.entity_uuid, CoordinatorAssign(user_uuid=coord.user_uuid))
        outsider_account = await _user(db, org, 31, "elsewhere")
        await db.execute(UserOrganization.__table__.delete().where(UserOrganization.user_id == outsider_account.id))
        await db.commit()

        job = await imp_svc.validate_upload(
            db, coord, entity.entity_uuid,
            _upload(_xlsx([["New", "", "new@x.com", "", "", "", ""], ["Else", "", "elsewhere@test.com", "", "", "", ""]])),
        )
        rows = {r.row_number: r for r in (await imp_svc.list_rows(db, coord, job.job_uuid)).items}
        assert rows[2].status == "valid" and rows[3].status == "invalid"

        with pytest.raises(HTTPException) as exc:
            await imp_svc.validate_upload(db, coord, other.entity_uuid, _upload(_xlsx([["a", "", "a@x.com", "", "", "", ""]])))
        assert exc.value.status_code == 403

        # A course the academy did not make available cannot be attached.
        course = Course(id=60, name="Cyber", description="", public=False, published=True, open_to_contributors=False,
                        org_id=org.id, course_uuid="course_cyber", creation_date="", update_date="")
        db.add(course)
        await db.commit()
        with pytest.raises(HTTPException) as exc:
            await imp_svc.commit_job(db, coord, job.job_uuid, ImportCommitOptions(resource_type="course", resource_uuid="course_cyber"))
        assert exc.value.status_code == 403
        await aud_svc.create_assignment(db, admin_user, AudienceAssignmentCreate(
            resource_type="course", resource_uuid="course_cyber", audience_type="entity", audience_uuid=entity.entity_uuid, mode="available"))
        with _patches():
            done = await imp_svc.commit_job(db, coord, job.job_uuid, ImportCommitOptions(resource_type="course", resource_uuid="course_cyber"))
        assert done.created_count == 1
        from src.security.rbac.rbac import check_usergroup_access

        new_user = (await db.execute(select(User).where(User.email == "new@x.com"))).scalar_one()
        assert await check_usergroup_access("course_cyber", new_user.id, db) is True

        await ent_svc.update_entity(db, admin_user, entity.entity_uuid, EntityUpdate(coordinator_permissions=CoordinatorPermissions(can_import_users=False)))
        with pytest.raises(HTTPException) as exc:
            await imp_svc.list_jobs(db, coord, entity.entity_uuid)
        assert exc.value.status_code == 403

    @pytest.mark.asyncio
    async def test_template_has_dropdowns(self, db, org, admin_user):
        entity, _group = await _setup(db, org, admin_user)
        content = await imp_svc.template_workbook(db, admin_user, entity.entity_uuid)
        wb = load_workbook(io.BytesIO(content))
        assert wb.sheetnames == ["Members", "Lists", "Instructions"]
        assert wb["Lists"].sheet_state == "hidden"
        assert [c.value for c in wb["Lists"]["A"]][1:] == ["Engineer"]
        assert len(wb["Members"].data_validations.dataValidation) == 2
        # The template's own headers parse back (an empty template has no rows).
        with pytest.raises(HTTPException) as exc:
            imp_svc.parse_rows("t.xlsx", content)
        assert exc.value.detail == "No rows to import"


class TestCoordinatorFollowUp:
    @pytest.mark.asyncio
    async def test_progress_only_covers_entity_learning_for_coordinators(self, db, org, admin_user, regular_user):
        entity, group = await _setup(db, org, admin_user)
        coord = await _user(db, org, 32, "coord2")
        await ent_svc.assign_coordinator(db, admin_user, entity.entity_uuid, CoordinatorAssign(user_uuid=coord.user_uuid))
        member = await ent_svc.add_member(
            db, admin_user, entity.entity_uuid,
            EntityMemberCreate(user_uuid=regular_user.user_uuid, group_uuids=[group.usergroup_uuid]),
        )
        c1 = Course(id=61, name="Entity course", description="", public=False, published=True, open_to_contributors=False, org_id=org.id, course_uuid="course_e", creation_date="", update_date="")
        c2 = Course(id=62, name="Private course", description="", public=True, published=True, open_to_contributors=False, org_id=org.id, course_uuid="course_p", creation_date="", update_date="")
        db.add_all([c1, c2])
        trail = Trail(org_id=org.id, user_id=regular_user.id, trail_uuid="trail_r")
        db.add(trail)
        await db.commit()
        for c, status in ((c1, "STATUS_COMPLETED"), (c2, "STATUS_IN_PROGRESS")):
            db.add(TrailRun(trail_id=trail.id, course_id=c.id, org_id=org.id, user_id=regular_user.id, status=status, creation_date=str(datetime.now()), update_date=str(datetime.now())))
        await db.commit()
        await aud_svc.create_assignment(db, admin_user, AudienceAssignmentCreate(
            resource_type="course", resource_uuid="course_e", audience_type="entity", audience_uuid=entity.entity_uuid, mode="available"))

        academy_view = await coord_svc.get_entity_progress(db, admin_user, entity.entity_uuid)
        row = next(m for m in academy_view.members if m.member_uuid == member.member_uuid)
        assert row.enrolled == 2 and row.completed == 1

        coord_view = await coord_svc.get_entity_progress(db, coord, entity.entity_uuid, group.usergroup_uuid)
        assert [m.member_uuid for m in coord_view.members] == [member.member_uuid]
        assert [c.course_uuid for c in coord_view.members[0].courses] == ["course_e"]
        assert coord_view.completion_rate == 100.0 and coord_view.courses[0].completed == 1

    @pytest.mark.asyncio
    async def test_instructor_invite_is_pending_and_hides_rates(self, db, org, admin_user):
        entity, _group = await _setup(db, org, admin_user)
        coord = await _user(db, org, 33, "coord3")
        await ent_svc.assign_coordinator(db, admin_user, entity.entity_uuid, CoordinatorAssign(user_uuid=coord.user_uuid))
        invite = EntityInstructorInvite(first_name="Sara", email="sara@ministry.gov.eg", specializations=["Networks"])
        with pytest.raises(HTTPException) as exc:
            await coord_svc.invite_instructor(db, coord, entity.entity_uuid, invite)
        assert exc.value.status_code == 403  # off by default
        await ent_svc.update_entity(db, admin_user, entity.entity_uuid, EntityUpdate(coordinator_permissions=CoordinatorPermissions(can_add_instructors=True)))
        with _patches():
            read = await coord_svc.invite_instructor(db, coord, entity.entity_uuid, invite)
        assert read.status == "pending_approval" and read.specializations == ["Networks"]
        assert "hourly_rate" not in read.model_dump() and "category" not in read.model_dump()
        instructor = (await db.execute(select(Instructor).where(Instructor.instructor_uuid == read.instructor_uuid))).scalar_one()
        assert instructor.entity_id is not None and instructor.hourly_rate is None
        with _patches(), pytest.raises(HTTPException) as exc:
            await coord_svc.invite_instructor(db, coord, entity.entity_uuid, invite)
        assert exc.value.status_code == 409
        assert [i.instructor_uuid for i in await coord_svc.list_entity_instructors(db, coord, entity.entity_uuid)] == [read.instructor_uuid]
